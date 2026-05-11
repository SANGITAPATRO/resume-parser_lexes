from __future__ import annotations

import io
import json
import logging
import os
import re
import threading
import traceback
import warnings
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeoutError
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional
from collections import defaultdict
import fitz                          # PyMuPDF
import geonamescache
import phonenumbers
import pdfplumber
import pytesseract
import requests
import spacy
from dateutil import parser as dateutil_parser
from docx import Document
from phonenumbers import PhoneNumberFormat, PhoneNumberMatcher
from PIL import Image
from skillNer.general_params import SKILL_DB
from skillNer.skill_extractor_class import SkillExtractor
from spacy.matcher import PhraseMatcher
from dotenv import load_dotenv
from constants import (
    _REAL_WORDS_ENDING_WITH_CONNECTIVE,
    _STATES,
    _SUMMARY_PROTECT,
    SECTION_KEYWORDS,
    EDUCATION_LEVELS,
    SKILL_CATEGORIES,
)

load_dotenv()
# ── Bootstrap ─────────────────────────────────────────────────────────────────

warnings.filterwarnings("ignore", message=r".*FontBBox.*")
warnings.filterwarnings("ignore", message=r".*empty vectors.*")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s — %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("resume_parser")

TESSERACT_CMD = (
    os.getenv('TESSERACT_PATH')
    if not __import__("os").environ.get("TESSERACT_PATH")
    else __import__("os").environ["TESSERACT_PATH"]
)
pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD


# ── Lazy-loaded singletons (thread-safe double-checked locking) ───────────────

_nlp: Optional[spacy.language.Language] = None
_skill_extractor: Optional[SkillExtractor] = None
_gc = geonamescache.GeonamesCache()

_NLP_LOCK   = threading.Lock()
_SKILL_LOCK = threading.Lock()


def get_nlp() -> spacy.language.Language:
    global _nlp
    if _nlp is not None:
        return _nlp
    with _NLP_LOCK:
        if _nlp is None:
            log.info("Loading spaCy model …")
            _nlp = spacy.load("en_core_web_lg")
    return _nlp


def get_skill_extractor() -> SkillExtractor:
    global _skill_extractor
    if _skill_extractor is not None:
        return _skill_extractor
    with _SKILL_LOCK:
        if _skill_extractor is None:
            log.info("Loading SkillExtractor …")
            _skill_extractor = SkillExtractor(get_nlp(), SKILL_DB, PhraseMatcher)
    return _skill_extractor


# ── Custom exceptions ─────────────────────────────────────────────────────────

class UnsupportedResumeFormat(Exception):
    """Raised when the file extension is not PDF / DOCX / DOC."""


class EmptyResumeError(Exception):
    """Raised when no text could be extracted from the file."""


# ── Compiled regex patterns ───────────────────────────────────────────────────

_EMAIL_RE    = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
_PINCODE_RE  = re.compile(r"\b\d{6}\b")
_LINKEDIN_RE = re.compile(r"linkedin\.com/in/[\w\-%.]+", re.IGNORECASE)
_GITHUB_RE   = re.compile(r"github\.com/[\w\-]+", re.IGNORECASE)

_LABELED_EMAIL_RE = re.compile(
    r"(?:"
    r"e[\s\-]?mail(?:\s*id)?|mail(?:\s*id)?|"
    r"alternate\s+(?:gmail|outlook|yahoo|email)(?:\s*id)?|"
    r"gmail(?:\s*id)?|outlook(?:\s*id)?|yahoo(?:\s*id)?|"
    r"contact\s*(?:email|mail)|personal\s*(?:email|mail)|official\s*(?:email|mail)"
    r")"
    r"\s*[:\-–]?\s*"
    r"([a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+)",
    re.IGNORECASE,
)

_MAILTO_RE = re.compile(
    r"mailto:\s*([a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+)",
    re.IGNORECASE,
)

_INLINE_EMAIL_LABEL_RE = re.compile(
    r"(?:email|e[\s\-]?mail|mail)\s*[:\-–]\s*"
    r"([a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+)",
    re.IGNORECASE,
)

_SPLIT_DOMAIN_PROVIDER_RE = re.compile(
    r"([a-zA-Z0-9_.+-]+@(?:gmail|yahoo|outlook|hotmail|rediffmail|icloud|"
    r"protonmail|live|zoho|ymail|msn|aol))\s*$",
    re.IGNORECASE,
)
_SPLIT_DOMAIN_TLD_RE = re.compile(
    r"^\s*\.(com|in|net|org|co\.in|co|edu|gov|me|info|io)\b",
    re.IGNORECASE,
)

_DOB_RE = re.compile(
    r"(?:date\s+of\s+birth|dob)\s*[:\-]\s*(\d{1,2}[./-]\d{1,2}[./-]\d{4})",
    re.IGNORECASE,
)

_DATE_RANGE_RE = re.compile(
    r"([A-Za-z]{3,9}\.?\s+\d{4}|\d{4})\s*[–\-]\s*"
    r"(Current|Present|[A-Za-z]{3,9}\.?\s+\d{4}|\d{4})",
    re.IGNORECASE,
)
_NUMERIC_DATE_RANGE_RE = re.compile(
    r"(\d{2}[/.-]\d{2}[/.-]\d{4})\s*[-–]\s*"
    r"(\d{2}[/.-]\d{2}[/.-]\d{4}|[Cc]urrent(?:ly)?|[Pp]resent)",
    re.IGNORECASE,
)
_YEAR_RANGE_RE = re.compile(
    r"\b(\d{4})\s*[-–]\s*(Present|Current|\d{4})\b",
    re.IGNORECASE,
)

_DATE_RE = re.compile(
    r"(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4}|\d{4})"
    r"\s*[-–—]\s*"
    r"(?:Current|Present|Ongoing|Till\s*date|Now|\d{4}"
    r"|(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4})",
    re.IGNORECASE,
)

_JOB_TITLE_RE = re.compile(
    r"^(?:Senior|Lead|Principal|Junior|Associate|Sr\.?|Jr\.?|Full[- ]Stack)?\s*"
    r"[A-Za-z][A-Za-z\s/&-]*"
    r"(?:Engineer|Developer|Analyst|Lead|Manager|Architect|Consultant|"
    r"Specialist|Director|Head|Coordinator|Administrator|Officer|Executive|"
    r"Support|Intern|Trainee|Associate|Programmer|Designer|Scientist|"
    r"Researcher|Strategist|Advisor)\b",
    re.IGNORECASE,
)

_COMPANY_SUFFIX_RE = re.compile(
    r"\b(Pvt\.?|Private|Ltd\.?|Limited|LLP|Inc\.?|Corp\.?|Corporation|LLC|"
    r"Solutions|Systems|Technologies|Group|International|India|Software|Services|"
    r"Consulting|Consultancy|Analytics|Innovations|Ventures|Holdings|"
    r"Global|Delivery)\b",
    re.IGNORECASE,
)

_ORG_HINTS_RE = re.compile(
    r"(.+?\b(?:college|school|university|institute|academy|vidyalaya|"
    r"polytechnic|campus|faculty|engineering|technology|management|"
    r"iit|nit|iiit|iim))\b",
    re.IGNORECASE,
)

_SCORE_RE = re.compile(
    r"(?:cgpa|gpa|percentage|marks|score|grade)[\s:]*([0-9]{1,3}(?:\.[0-9]{1,2})?)"
    r"|([0-9]{1,3}(?:\.[0-9]{1,2})?)\s*/\s*(10|100)",
    re.IGNORECASE,
)

_BULLET_RE      = re.compile(r"^[\s•\-\*✓➔→◦·▪▸◆●]+")
_SEPARATOR_RE   = re.compile(r"^[=\-_*#~]{5,}\s*$")
_PAREN_YEAR_RE  = re.compile(r"\((\d{4})\s*[-–]?\s*(\d{4}|[Pp]resent)?\)")
_SPACED_CAPS_RE = re.compile(r"\b[A-Z](?:\s+[A-Z]){2,}\b")
_BIG_GAP_RE     = re.compile(r"\s{6,}")
_EMOJI_RE       = re.compile(
    r"[\U00010000-\U0010ffff\u2600-\u26ff\u2700-\u27bf\U0001f300-\U0001f9ff]+",
    re.UNICODE,
)
_SKILL_STOPWORDS = frozenset({
    "and", "or", "the", "of", "in", "for", "to", "with",
    "using", "based", "driven", "oriented",
})
_NOISE_PATTERNS = [
    re.compile(r"^(present|current|ongoing|till|now|designed|enforced?|collaborated?)\b", re.I),
    re.compile(r"^[a-z]\s+[a-z]$", re.I),
    re.compile(r"^\d+"),
    re.compile(r"[#@/\\]{2,}"),
]

_SKILL_ACTION_VERB_RE = re.compile(
    r"^(conducted|delivered|optimized|managed|oversaw|assisted|drove|built|led|"
    r"created|developed|designed|implemented|executed|ensured|provided|handled|"
    r"collaborated|worked|performed|identified|supported|maintained|coordinated|"
    r"guided|prepared|utilized|engaged|fulfilled|closed|scaled|partnered|while|"
    r"i\s+was|i\s+have|however|including|contributing|driving|independently|"
    r"personally)",
    re.IGNORECASE,
)
_SKILL_META_LABEL_RE = re.compile(
    r"^(role\s*/?\s*designation|credential\s*id|issuing\s*authority|"
    r"issued\s*date|responsibilities|achievements\s*(?:&|and)\s*awards)\s*[:\-]?",
    re.IGNORECASE,
)
_SKILL_FINANCIAL_RE = re.compile(
    r"(inr|usd|sgd)|s\$|₹|revenue|turnover|salary|budget|p&l|\d+[kKcC][rR]?",
    re.IGNORECASE,
)
_SKILL_GLUED_RE  = re.compile(r"[a-z]{3,}(?:to|and|in|or|of)\s+[a-z]")
_SKILL_TECH_KW_RE = re.compile(
    r"(api|sql|aws|azure|gcp|python|java|react|angular|node|docker|"
    r"kubernetes|git|rest|json|xml|html|css|javascript|typescript|"
    r"testing|etl|bi|crm|erp|sap|devops|agile|scrum|net|tableau|"
    r"informatica|vertica|nifi|jira|jenkins|selenium|kafka|spark|"
    r"jquery|vb|c#|xaml|wpf|winforms|mvc|asp)",
    re.IGNORECASE,
)
_SKILL_KNOWN_COUNTRIES = frozenset({
    "australia", "uk", "usa", "canada", "singapore", "india", "malaysia",
    "indonesia", "philippines", "russia", "brazil", "china", "japan", "germany",
    "france", "dubai", "uae",
})
_SKILL_KNOWN_COLLEGES = frozenset({"nit", "iit", "bits", "vit", "iim", "iiit", "ignou"})
_SKILL_GENERIC_WORDS = frozenset({
    "however", "etc", "process", "sessions", "transition", "basis", "update",
    "bug", "writing", "office", "reports", "partner", "develop", "motivated",
    "security", "compliance", "delivery", "framework", "incident", "management",
    "automation", "code", "debugging", "prototyping", "milestones", "deliverables",
    "dashboard", "transformation", "parse",
})
_SKILL_TECH_WHITELIST = frozenset({
    "sql", "aws", "gcp", "azure", "git", "api", "sdk", "etl", "bi", "ml", "ai",
    "rest", "soap", "xml", "json", "css", "html", "php", "c#", "r", "go",
    "jira", "ci", "cd", "qa", "ux", "ui", "erp", "crm", "sap", "ios", "nlp",
    "dba", "sre", "devops", "scrum", "agile", "linux", "bash", "rust", "nifi",
    "jquery", "angular", "react", "vue", "redux", "webpack",
})
_SKILL_COMPANY_SUFFIX_RE = re.compile(
    r"(pvt\.?\s*ltd\.?|limited|llc|inc\.?|corp\.?|private\s+limited|"
    r"solutions\s+private|technologies\s+private|systems\s+private)",
    re.IGNORECASE,
)


def _is_skill_noise_extended(skill: str) -> bool:
    s     = skill.strip()
    lower = s.lower()
    words = s.split()

    if len(s) < 2:
        return True
    if s[0].islower() and lower not in _SKILL_TECH_WHITELIST:
        if len(words[0]) <= 3 and words[0].islower():
            return True
        if len(words) == 1 and lower in _SKILL_GENERIC_WORDS:
            return True
    if _SKILL_GLUED_RE.search(s) and not _SKILL_TECH_KW_RE.search(s):
        return True
    if s[0] in ("/", "\\"):
        return True
    if _SKILL_ACTION_VERB_RE.match(s):
        return True
    if _SKILL_META_LABEL_RE.match(s):
        return True
    if _SKILL_FINANCIAL_RE.search(s):
        return True
    if _SKILL_COMPANY_SUFFIX_RE.search(s):
        return True
    upper_words = [w for w in words if w.isupper() and len(w) > 1]
    if len(upper_words) >= 2 and len(words) <= 4 and not _SKILL_TECH_KW_RE.search(s):
        return True
    if lower in _SKILL_KNOWN_COUNTRIES and len(words) == 1:
        return True
    if lower in _SKILL_KNOWN_COLLEGES and len(words) == 1:
        return True
    if len(words) == 1 and lower in _SKILL_GENERIC_WORDS:
        return True
    if len(words) >= 6 and not _SKILL_TECH_KW_RE.search(s):
        return True
    return False


_KNOWN_FIXES = {
    "c #":        "C#",
    "c#":         "C#",
    "asp net":    "ASP.NET",
    "node js":    "Node.js",
    "vue js":     "Vue.js",
    "react js":   "React.js",
    "dot net":    ".NET",
    ".net core":  ".NET Core",
    "power bi":   "Power BI",
    "my sql":     "MySQL",
}


# ── Data models ───────────────────────────────────────────────────────────────

@dataclass
class DateField:
    date: Optional[str] = None
    is_current: bool = False
    day: Optional[int] = None
    month: Optional[int] = None
    year: Optional[int] = None


@dataclass
class WorkExperience:
    job_title: str = ""
    organization: str = ""
    date_range: str = ""
    start_date: Optional[DateField] = None
    end_date: Optional[DateField] = None
    description: str = ""
    location: str = ""


@dataclass
class Education:
    organization: str = ""
    degree: str = ""
    field_of_study: str = ""
    level: str = ""
    start_year: Optional[int] = None
    end_year: Optional[int | str] = None
    grade: Optional[dict[str, Any]] = None


@dataclass
class Project:
    title: str = ""
    client: str = ""
    role: str = ""
    description: str = ""
    technologies: list[str] = field(default_factory=list)
    dates: Optional[dict[str, Any]] = None
    project_type: str = "professional"


# ── Text-extraction helpers ───────────────────────────────────────────────────

def _extract_pdf_pdfplumber(path: Path) -> str:
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text(layout=True, x_tolerance=2, y_tolerance=2)
            if text:
                pages.append(text)
    return "\n".join(pages)


def _extract_pdf_pymupdf(path: Path) -> str:
    doc = fitz.open(path)
    pages = [page.get_text("text") for page in doc]
    doc.close()
    return "\n".join(pages)


def _extract_pdf_ocr(path: Path) -> str:
    doc = fitz.open(path)
    results: list[str] = []
    for page in doc:
        pix = page.get_pixmap(dpi=300)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        results.append(pytesseract.image_to_string(img))
    doc.close()
    return "\n".join(results)


def extract_pdf_text(path: Path) -> str:
    for extractor, label in [
        (_extract_pdf_pdfplumber, "pdfplumber"),
        (_extract_pdf_pymupdf,    "PyMuPDF"),
        (_extract_pdf_ocr,        "OCR"),
    ]:
        try:
            text = extractor(path)
            if len(text.strip()) >= 50:
                log.debug("Extracted via %s", label)
                return text
        except Exception as exc:
            log.warning("%s failed: %s", label, exc)
    return ""


def extract_docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


# ── Layout helpers ────────────────────────────────────────────────────────────

def detect_two_column_pdf(path: Path) -> bool:
    try:
        with pdfplumber.open(path) as pdf:
            page  = pdf.pages[0]
            words = page.extract_words()
        if not words:
            return False
        mid   = page.width / 2
        left  = sum(1 for w in words if w["x0"] < mid)
        right = sum(1 for w in words if w["x0"] >= mid)
        if right < 30:
            return False
        ratio = min(left, right) / max(left, right) if max(left, right) else 0
        return ratio > 0.35
    except Exception as exc:
        log.warning("Column detection failed: %s", exc)
        return False


# ── Text-normalisation utilities ──────────────────────────────────────────────

def _normalize_text(text: str) -> str:
    text = text.replace("\r", "").replace("\t", " ")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = _SPACED_CAPS_RE.sub(lambda m: m.group(0).replace(" ", ""), text)
    return text.strip()


def _strip_emoji(text: str) -> str:
    return _EMOJI_RE.sub("", text).strip()


def _clean_bullet(line: str) -> str:
    return _BULLET_RE.sub("", line).strip()


def _is_bullet_line(line: str) -> bool:
    return bool(_BULLET_RE.match(line.strip()))


def _is_separator(line: str) -> bool:
    return bool(_SEPARATOR_RE.match(line.strip()))


def _lines(text: str) -> list[str]:
    return [ln.strip() for ln in text.splitlines() if ln.strip()]


# ── Section segmentation ──────────────────────────────────────────────────────

def _detect_section_header(line: str) -> Optional[str]:
    normalized = _strip_emoji(line).lower().strip().rstrip(":").strip()
    for section, keywords in SECTION_KEYWORDS.items():
        if normalized == section.lower() or normalized in keywords:
            return section
    return None


def _is_section_header_heuristic(line: str) -> bool:
    stripped = _strip_emoji(line).strip()
    if len(stripped) > 60 or len(stripped) < 3:
        return False
    return stripped.isupper() and not any(ch.isdigit() for ch in stripped)


def segment_resume(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {k: [] for k in SECTION_KEYWORDS}
    current = "others"

    for line in text.splitlines():
        clean = _strip_emoji(line)
        header = _detect_section_header(clean)
        if header:
            current = header
            continue
        if _is_section_header_heuristic(clean):
            candidate = _detect_section_header(clean)
            if candidate:
                current = candidate
                continue
        if line.strip():
            sections[current].append(line)

    return {k: "\n".join(v).strip() for k, v in sections.items() if "".join(v).strip()}


# ── Field parsers ─────────────────────────────────────────────────────────────

def parse_candidate_name(text: str) -> dict[str, Optional[str]]:
    _EXCLUDE = re.compile(
        r"\d|@|linkedin|github|http|www|phone|email|mobile|address|resume|cv",
        re.IGNORECASE,
    )
    _NAME_TOKEN_RE = re.compile(r"^[A-Za-z]([a-zA-Z\'\-\.]*[a-zA-Z])?$")
    _PAREN_ALIAS_RE = re.compile(r"\s*\([^)]{1,30}\)\s*")
    _ROLE_GUARD_RE = re.compile(
        r"\b(engineer|developer|analyst|manager|designer|architect|consultant|"
        r"director|officer|specialist|lead|intern|trainee|certif)\b",
        re.IGNORECASE,
    )
    _SECTION_HDR_RE = re.compile(
        r"^(SKILLS|EDUCATION|EXPERIENCE|SUMMARY|OBJECTIVE|PROFILE|CONTACT|"
        r"PROJECTS|CERTIFICATIONS|ACHIEVEMENTS|LANGUAGES|DECLARATION|REFERENCES|"
        r"INTERNSHIPS|ACTIVITIES|INTERESTS|PUBLICATIONS|VOLUNTEERING|PASSPORT|"
        r"ABOUT|AWARDS|STRENGTHS|TOOLS|QUALIFICATIONS)$",
        re.IGNORECASE,
    )

    for line in _lines(text)[:20]:
        clean_line = _PAREN_ALIAS_RE.sub(" ", line).strip()
        if not clean_line or _EXCLUDE.search(clean_line):
            continue
        if _ROLE_GUARD_RE.search(clean_line):
            continue
        if _SECTION_HDR_RE.match(clean_line):
            continue
        parts = clean_line.split()
        if 2 <= len(parts) <= 5:
            if all(_NAME_TOKEN_RE.match(p) for p in parts):
                return {
                    "first_name":  parts[0],
                    "middle_name": " ".join(parts[1:-1]) if len(parts) > 2 else None,
                    "family_name": parts[-1],
                    "full_name":   " ".join(parts),
                }
        elif len(parts) == 1:
            token = parts[0]
            if (
                len(token) >= 4
                and token.isalpha()
                and (token.isupper() or token.istitle())
                and not _SECTION_HDR_RE.match(token)
            ):
                return {
                    "first_name":  token,
                    "middle_name": None,
                    "family_name": None,
                    "full_name":   token,
                }

    return {}


# ── Email helpers ─────────────────────────────────────────────────────────────

def _repair_split_domain_emails(text: str) -> list[str]:
    results: list[str] = []
    lines = text.splitlines()
    for i, line in enumerate(lines):
        m_provider = _SPLIT_DOMAIN_PROVIDER_RE.search(line)
        if not m_provider:
            continue
        next_line = ""
        for j in range(i + 1, min(i + 4, len(lines))):
            candidate = lines[j].strip()
            if candidate:
                next_line = candidate
                break
        if not next_line:
            continue
        m_tld = _SPLIT_DOMAIN_TLD_RE.match(next_line)
        if not m_tld:
            continue
        fragment  = m_provider.group(1).strip()
        tld       = m_tld.group(0).strip()
        full_email = (fragment + tld).lower()
        if re.match(r"[a-z0-9_.+-]+@[a-z0-9-]+\.[a-z]{2,}", full_email):
            results.append(full_email)
    return results


def _score_email(email: str) -> float:
    score = 0.0
    if email.count("@") == 1:
        score += 0.4
        local, domain = email.split("@")
        if len(local) >= 6:
            score += 0.2
        else:
            score -= 0.3
        if re.search(r"\.(com|in|net|org|io|co\.in|edu|gov|me|info)$", domain, re.I):
            score += 0.2
        if ".." not in email:
            score += 0.1
        if not re.search(r"[\s,;:]", email):
            score += 0.1
    return round(score, 2)


def _collect_email_candidates(text: str) -> list[tuple[str, str, float]]:
    seen:       set[str]                     = set()
    candidates: list[tuple[str, str, float]] = []

    def _add(email: str, source: str) -> None:
        e = email.strip().lower()
        if not e or e in seen:
            return
        seen.add(e)
        candidates.append((e, source, _score_email(e)))

    for m in _LABELED_EMAIL_RE.finditer(text):
        _add(m.group(1), "labeled_field")
    for m in _INLINE_EMAIL_LABEL_RE.finditer(text):
        _add(m.group(1), "inline_label")
    for m in _MAILTO_RE.finditer(text):
        _add(m.group(1), "mailto_uri")
    for email in _repair_split_domain_emails(text):
        _add(email, "split_domain")

    text_stitched = re.sub(
        r"(@(?:gmail|yahoo|outlook|hotmail|rediffmail|icloud|protonmail|"
        r"live|zoho|ymail|msn|aol))\s*\n\s*(\.[a-zA-Z]{2,6})\b",
        r"\1\2", text, flags=re.IGNORECASE,
    )
    text_joined = re.sub(r"\s*\n\s*", " ", text_stitched)
    for m in _EMAIL_RE.finditer(text_joined):
        _add(m.group(0), "raw_regex")

    return candidates


def _select_best_emails(candidates: list[tuple[str, str, float]]) -> list[str]:
    SOURCE_PRIORITY = {
        "mailto_uri":    0,
        "labeled_field": 1,
        "inline_label":  2,
        "split_domain":  3,
        "raw_regex":     4,
    }
    valid = [(e, src, sc) for e, src, sc in candidates if sc >= 0.6]
    if not valid and candidates:
        valid = list(candidates)
    valid.sort(key=lambda x: (-x[2], SOURCE_PRIORITY.get(x[1], 9)))
    seen_local: set[str] = set()
    result: list[str] = []
    for email, _src, _sc in valid:
        local = email.split("@")[0]
        if local not in seen_local:
            seen_local.add(local)
            result.append(email)
    return result


def parse_contact(text: str) -> dict[str, Any]:
    candidates  = _collect_email_candidates(text)
    best_emails = _select_best_emails(candidates)
    primary_email = best_emails[0] if best_emails else ""

    phones: list[str] = sorted({
        phonenumbers.format_number(m.number, PhoneNumberFormat.E164)
        for m in PhoneNumberMatcher(text, "IN")
    })

    linkedin = github = ""
    lm = _LINKEDIN_RE.search(text)
    if lm:
        linkedin = "https://" + lm.group(0).lower()
    gm = _GITHUB_RE.search(text)
    if gm:
        github = "https://" + gm.group(0).lower()

    other_urls = re.findall(
        r"https?://(?!linkedin|github)[^\s,>\"]+", text, re.IGNORECASE
    )

    return {
        "email":      primary_email,
        "phone":      phones,
        "linkedin":   linkedin,
        "github":     github,
        "other_urls": other_urls,
        "all_emails": best_emails,
        "email_candidates": [
            {"email": e, "source": s, "score": sc}
            for e, s, sc in candidates
        ],
    }


# ── Date parsing ──────────────────────────────────────────────────────────────

_CURRENT_TERMS = frozenset(
    {"present", "current", "currently", "ongoing", "till date", "till now", "now"}
)


def parse_date(raw: str) -> DateField:
    norm = raw.strip().lower()
    if norm in _CURRENT_TERMS or "present" in norm:
        return DateField(is_current=True)
    for fmt in ("%B %Y", "%b %Y", "%b. %Y"):
        try:
            dt = datetime.strptime(raw.strip().title(), fmt)
            return DateField(date=dt.date().isoformat(), month=dt.month, year=dt.year)
        except ValueError:
            pass
    if raw.strip().isdigit() and len(raw.strip()) == 4:
        return DateField(year=int(raw.strip()))
    nm = re.match(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", raw.strip())
    if nm:
        try:
            d = date(int(nm.group(3)), int(nm.group(2)), int(nm.group(1)))
            return DateField(date=d.isoformat(), day=d.day, month=d.month, year=d.year)
        except ValueError:
            pass
    try:
        dt = dateutil_parser.parse(raw.strip(), default=datetime(2000, 1, 1))
        return DateField(date=dt.date().isoformat(), month=dt.month, year=dt.year)
    except Exception:
        pass
    return DateField()


# ── Experience parsing ────────────────────────────────────────────────────────

def _is_job_title(line: str) -> bool:
    clean = _strip_emoji(line).strip()
    if _is_bullet_line(clean) or len(clean) > 120:
        return False
    return bool(_JOB_TITLE_RE.match(clean))


def _extract_any_date_range(line: str) -> Optional[tuple[str, str, str]]:
    m = _DATE_RE.search(line)
    if m:
        raw = m.group(0)
        r2  = _DATE_RANGE_RE.search(raw)
        if r2:
            return raw, r2.group(1), r2.group(2)
        return raw, "", ""
    m2 = _NUMERIC_DATE_RANGE_RE.search(line)
    if m2:
        return m2.group(0), m2.group(1), m2.group(2)
    m3 = _YEAR_RANGE_RE.search(line)
    if m3:
        return m3.group(0), m3.group(1), m3.group(2)
    return None


def _try_split_inline_exp(line: str) -> Optional[tuple[str, str, str]]:
    clean = _strip_emoji(line).strip()
    m = re.match(
        r"^([A-Za-z][A-Za-z\s/&-]+?"
        r"(?:Engineer|Developer|Analyst|Lead|Manager|Architect|Consultant|"
        r"Specialist|Director|Officer|Executive|Support|Intern|Trainee|"
        r"Programmer|Designer|Scientist|Researcher|Strategist))"
        r"\s*[-–]\s*"
        r"(.+?)\s+"
        r"(\d{4}\s*[-–]\s*(?:Present|Current|\d{4}))\s*$",
        clean, re.IGNORECASE,
    )
    if m:
        return m.group(1).strip(), m.group(2).strip(), m.group(3).strip()
    return None


def parse_experience(text: str) -> list[dict[str, Any]]:
    if not text.strip():
        return []
    lines      = [_strip_emoji(ln).strip() for ln in text.splitlines() if _strip_emoji(ln).strip()]
    experiences: list[WorkExperience] = []
    current: Optional[WorkExperience] = None

    def _commit() -> None:
        if current and current.job_title:
            experiences.append(current)

    for line in lines:
        if len(line) < 3 or _is_separator(line):
            continue
        inline = _try_split_inline_exp(line)
        if inline:
            _commit()
            title, org, date_str = inline
            current = WorkExperience(job_title=title, organization=org)
            dr = _extract_any_date_range(date_str)
            if dr:
                current.date_range = dr[0]
                if dr[1]: current.start_date = parse_date(dr[1])
                if dr[2]: current.end_date   = parse_date(dr[2])
            continue
        if _is_job_title(line) and not _is_bullet_line(line):
            _commit()
            current = WorkExperience(job_title=line)
            continue
        if current is None:
            continue
        dr = _extract_any_date_range(line)
        if dr:
            raw_date, start_str, end_str = dr
            current.date_range = raw_date
            if start_str: current.start_date = parse_date(start_str)
            if end_str:   current.end_date   = parse_date(end_str)
            date_pos = line.find(raw_date)
            if date_pos > 0:
                pre = line[:date_pos].strip()
                if pre and not current.organization and not _is_bullet_line(pre) and len(pre) < 80:
                    current.organization = pre
            continue
        if (
            not current.organization
            and _COMPANY_SUFFIX_RE.search(line)
            and not _is_bullet_line(line)
            and len(line) < 80
        ):
            current.organization = line
            continue
        desc_line = _clean_bullet(line)
        if desc_line:
            current.description = (
                (current.description + "\n" + desc_line).strip()
                if current.description else desc_line
            )
    _commit()
    return [_experience_to_dict(e) for e in experiences]


def _experience_to_dict(e: WorkExperience) -> dict[str, Any]:
    return {
        "work_experience_job_title":    e.job_title,
        "work_experience_organization": e.organization,
        "work_experience_date_range":   e.date_range,
        "work_experience_start_date":   asdict(e.start_date) if e.start_date else None,
        "work_experience_end_date":     asdict(e.end_date)   if e.end_date   else None,
        "work_experience_description":  e.description,
        "work_experience_location":     e.location,
    }


# ── H-07 FIX: helper functions for interval-based experience calculation ──────

_EXP_MENTION_RE = re.compile(
    r"(?:(?:with|having|over|more\s+than|around|about|i\s+have|of)\s+)?"
    r"(\d+(?:\.\d+)?)\s*\+?\s*"
    r"(?:years?|yrs?)"
    r"(?:\s+of\s+(?:experience|exp\.?|professional\s+experience|"
    r"industry\s+experience|relevant\s+experience|work\s+experience))?",
    re.IGNORECASE,
)


def _extract_mentioned_years(text: str) -> Optional[float]:
    values = [
        float(m.group(1))
        for m in _EXP_MENTION_RE.finditer(text)
        if 0.5 <= float(m.group(1)) <= 60
    ]
    return max(values) if values else None


def _to_month_index(year: int, month: int) -> int:
    return year * 12 + month


def _merge_month_intervals(intervals: list[tuple[int, int]]) -> list[tuple[int, int]]:
    if not intervals:
        return []
    intervals = sorted(intervals)
    merged: list[tuple[int, int]] = [intervals[0]]
    for s, e in intervals[1:]:
        if s <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], e))
        else:
            merged.append((s, e))
    return merged


def calculate_total_experience(
    exp_text: str,
    summary_text: str = "",
) -> float:
    now = datetime.now()
    raw_intervals: list[tuple[int, int]] = []

    def _end_month_index(end_str: str) -> int:
        if end_str.lower().strip() in _CURRENT_TERMS:
            return _to_month_index(now.year, now.month)
        try:
            dt = dateutil_parser.parse(end_str)
            return _to_month_index(dt.year, dt.month)
        except Exception:
            return _to_month_index(now.year, now.month)

    for start_str, end_str in _DATE_RANGE_RE.findall(exp_text):
        try:
            start = dateutil_parser.parse(start_str)
            si    = _to_month_index(start.year, start.month)
            ei    = _end_month_index(end_str)
            if ei > si:
                raw_intervals.append((si, ei))
        except Exception:
            continue

    for start_str, end_str in _NUMERIC_DATE_RANGE_RE.findall(exp_text):
        try:
            start = dateutil_parser.parse(start_str, dayfirst=True)
            si    = _to_month_index(start.year, start.month)
            ei    = _end_month_index(end_str)
            if ei > si:
                raw_intervals.append((si, ei))
        except Exception:
            continue

    if not raw_intervals:
        for start_str, end_str in _YEAR_RANGE_RE.findall(exp_text):
            try:
                sy = int(start_str)
                ey = now.year if end_str.lower() in _CURRENT_TERMS else int(end_str)
                si = _to_month_index(sy, 1)
                ei = _to_month_index(ey, 12)
                if ei > si:
                    raw_intervals.append((si, ei))
            except Exception:
                continue

    merged       = _merge_month_intervals(raw_intervals)
    total_months = sum(e - s for s, e in merged)
    if total_months > 0:
        return round(total_months / 12, 1)

    combined_text = exp_text + "\n" + summary_text
    mentioned = _extract_mentioned_years(combined_text)
    return mentioned if mentioned is not None else 0.0


# ══════════════════════════════════════════════════════════════════════════════
# ── PATCH: Current Position & Experience Extraction ───────────────────────────
# ══════════════════════════════════════════════════════════════════════════════

# Matches the most common Indian-resume summary opener, e.g.:
#   "Senior Software Engineer with 8+ years of experience"
#   "Results-oriented Team Lead with over 12 years"
#   "Dynamic Project Manager | 7+ Years of Experience"
#   "Experienced Data Scientist having 4.2+ yrs"
#   "Tech Lead – 10 years of experience in cloud"
_SUMMARY_POSITION_RE = re.compile(
    r"^(?:(?:experienced|results[\s-]oriented|dynamic|motivated|seasoned|"
    r"dedicated|strategic|accomplished|passionate|proactive|detail[\s-]oriented|"
    r"innovative|skilled|qualified|certified|self[\s-]motivated)\s+)?"
    r"((?:(?:senior|lead|principal|junior|associate|sr\.?|jr\.?|"
    r"chief|head\s+of|vp\s+of|avp|deputy|assistant|"
    r"full[\s-]stack)\s+)?"
    r"[A-Z][A-Za-z\s/&\-]{0,60}?"
    r"(?:engineer|developer|analyst|lead|manager|architect|consultant|"
    r"specialist|director|head|officer|executive|scientist|"
    r"researcher|designer|strategist|advisor|technician|"
    r"programmer|administrator|coordinator|intern|trainee|associate))"
    r"\s*(?:with|having|\u2013|\u2014|[|,;-])?\s*"
    r"(?:over|more\s+than|around|about|approximately|nearly|almost)?\s*"
    r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)"
    r"(?:\s+of\s+(?:experience|exp\.?|professional|industry|relevant|work))?",
    re.IGNORECASE,
)

# Catches pipe/dash-separated patterns:
#   "Project Manager | 7+ Years of Experience"
#   "Tech Lead – 10 years in cloud"
_SUMMARY_PIPE_RE = re.compile(
    r"((?:senior|lead|principal|junior|associate|sr\.?|jr\.?|"
    r"chief|head\s+of|vp\s+of|avp|deputy|assistant|"
    r"full[\s-]stack)\s+)?"
    r"([A-Z][A-Za-z\s/&\-]{2,60}?"
    r"(?:engineer|developer|analyst|lead|manager|architect|consultant|"
    r"specialist|director|head|officer|executive|scientist|"
    r"researcher|designer|strategist|advisor|technician|"
    r"programmer|administrator|coordinator|intern|trainee|associate))"
    r"\s*[|\u2013\u2014-]\s*"
    r"(?:over|more\s+than|around|about|approximately)?\s*"
    r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)",
    re.IGNORECASE,
)

# Looser fallback: title keyword anywhere in first block followed by year mention
_TITLE_THEN_YEARS_RE = re.compile(
    r"(?P<title>"
    r"(?:senior|lead|principal|jr\.?|sr\.?|associate|chief|deputy|assistant|"
    r"full[\s-]stack\s+)?"
    r"[A-Za-z][A-Za-z\s/&\-]{2,55}"
    r"(?:engineer|developer|analyst|lead|manager|architect|consultant|"
    r"specialist|director|officer|executive|scientist|researcher|"
    r"designer|strategist|advisor|programmer|administrator|coordinator))"
    r"[^\n.!?]{0,80}?"
    r"(?:with|having|\u2013|\u2014|[|])?\s*"
    r"(?:over|more\s+than|around|about|approximately)?\s*"
    r"(?P<years>\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)",
    re.IGNORECASE,
)

# Guards: skip section headers and duty/action sentences
_POSITION_SECTION_GUARD_RE = re.compile(
    r"^(SKILLS|EDUCATION|EXPERIENCE|WORK\s+EXPERIENCE|SUMMARY|OBJECTIVE|"
    r"PROFILE|CONTACT|PROJECTS|CERTIFICATIONS|ACHIEVEMENTS|LANGUAGES|"
    r"DECLARATION|REFERENCES|INTERNSHIPS|ACTIVITIES|INTERESTS|"
    r"PUBLICATIONS|VOLUNTEERING|ABOUT|AWARDS|STRENGTHS|TOOLS|"
    r"QUALIFICATIONS|CAREER\s+OBJECTIVE|PROFESSIONAL\s+SUMMARY|"
    r"CAREER\s+SUMMARY|PROFILE\s+SUMMARY)$",
    re.IGNORECASE,
)
_POSITION_DUTY_GUARD_RE = re.compile(
    r"\b(responsible|worked|developed|built|maintained|supported|"
    r"handled|managed\s+team|led\s+team|implemented|deployed|"
    r"collaborated|participated|ensured|achieved|delivered|provided)\b",
    re.IGNORECASE,
)

_SUMMARY_SCAN_LINES = 15

# Terms that indicate an open / current role in work_experience end_date
_CURRENT_END_TERMS = frozenset(
    {"present", "current", "currently", "ongoing",
     "till date", "till now", "now", "date"}
)

# Adjective / filler openers that Indian candidates write before their title.
# These must be stripped from the extracted title — they are NOT part of the
# job designation.  Pattern is applied to the raw matched group.
# Examples that must be cleaned:
#   "Results-driven Executive Manager"  → "Executive Manager"
#   "Results-oriented Technical Lead"   → "Technical Lead"
#   "Dynamic Project Manager"           → "Project Manager"
#   "Experienced Senior Developer"      → "Senior Developer"
#   "Highly motivated Software Engineer"→ "Software Engineer"
_TITLE_ADJECTIVE_PREFIX_RE = re.compile(
    r"^(?:"
    # single-word openers
    r"experienced|dynamic|motivated|seasoned|dedicated|strategic|"
    r"accomplished|passionate|proactive|innovative|skilled|qualified|"
    r"certified|driven|oriented|focused|committed|enthusiastic|"
    r"creative|analytical|diligent|hardworking|versatile|resourceful|"
    r"competent|proficient|talented|ambitious|goal[\s-]oriented|"
    r"result[\s-]oriented|results[\s-]driven|results[\s-]oriented|"
    r"detail[\s-]oriented|self[\s-]motivated|self[\s-]driven|"
    r"customer[\s-]focused|customer[\s-]centric|"
    r"highly\s+motivated|highly\s+skilled|highly\s+experienced|"
    r"well[\s-]experienced|well[\s-]versed|"
    # two-word openers
    r"result[s]?[\s-]driven|result[s]?[\s-]focused|"
    r"process[\s-]oriented|solution[\s-]oriented|"
    r"quality[\s-]driven|data[\s-]driven|"
    r"team[\s-]oriented|people[\s-]oriented"
    r")"
    r"[\s,;:-]+",          # trailing separator after the opener word(s)
    re.IGNORECASE,
)


def _clean_position_title(title: str) -> str:
    """
    Strip leading adjective/filler openers from a matched position title.
    Apply repeatedly (up to 3 passes) to handle stacked openers like
    'Highly motivated Results-driven Senior Engineer'.
    """
    for _ in range(3):
        cleaned = _TITLE_ADJECTIVE_PREFIX_RE.sub("", title).strip().strip(".,;:-").strip()
        if cleaned == title:
            break
        title = cleaned
    # Normalise internal whitespace
    return re.sub(r"\s{2,}", " ", title).strip()


def _build_position_result(
    title: str,
    years_raw: str,
    source: str,
    original_line: str,
) -> dict[str, Any]:
    """Build the standardised current_position result dict."""
    # Strip adjective openers BEFORE storing the title
    title = _clean_position_title(title)

    # Preserve '+' if it appeared directly after the digit in the original text
    search_start = original_line.find(years_raw)
    after = original_line[search_start: search_start + len(years_raw) + 2]
    has_plus = "+" in after
    years_display = f"{years_raw}{'+'  if has_plus else ''} years"
    try:
        years_value = float(years_raw)
    except (ValueError, TypeError):
        years_value = None
    return {
        "current_position":          title,
        "years_of_experience":       years_display,
        "years_of_experience_value": years_value,
        "source":                    source,
    }


def _scan_summary_for_position(text: str) -> Optional[dict[str, Any]]:
    """Tier 1 — scan the summary/profile/about section for position + years."""
    lines = [_strip_emoji(ln).strip() for ln in text.splitlines() if _strip_emoji(ln).strip()]

    for line in lines[:_SUMMARY_SCAN_LINES]:
        if _POSITION_SECTION_GUARD_RE.match(line):
            continue

        # Strategy A — main combined pattern
        m = _SUMMARY_POSITION_RE.search(line)
        if m:
            title = m.group(1).strip().strip(".,;:-")
            if not _POSITION_DUTY_GUARD_RE.search(title) and 3 <= len(title) <= 80:
                return _build_position_result(title, m.group(2), "summary", line)

        # Strategy B — pipe / dash separated
        m2 = _SUMMARY_PIPE_RE.search(line)
        if m2:
            prefix = (m2.group(1) or "").strip()
            base   = m2.group(2).strip().strip(".,;:-")
            title  = (prefix + " " + base).strip() if prefix else base
            if not _POSITION_DUTY_GUARD_RE.search(title) and 3 <= len(title) <= 80:
                return _build_position_result(title, m2.group(3), "summary", line)

    # Strategy C — looser full-block scan
    first_block = " ".join(lines[:_SUMMARY_SCAN_LINES])
    m3 = _TITLE_THEN_YEARS_RE.search(first_block)
    if m3:
        title = _clean_position_title(m3.group("title").strip().strip(".,;:-"))
        if not _POSITION_DUTY_GUARD_RE.search(title) and 3 <= len(title) <= 80:
            return _build_position_result(title, m3.group("years"), "summary", first_block)

    return None


def _is_current_role(exp_entry: dict[str, Any]) -> bool:
    """Return True if this work_experience entry appears to be an open/current role."""
    end_date = exp_entry.get("work_experience_end_date")
    if end_date:
        if end_date.get("is_current"):
            return True
        if not end_date.get("year") and not end_date.get("date"):
            return True
    else:
        return True  # no end_date key at all → still in role

    date_range = (exp_entry.get("work_experience_date_range") or "").lower()
    return any(term in date_range for term in _CURRENT_END_TERMS)


def _extract_from_work_experience(
    work_experience: list[dict[str, Any]],
) -> Optional[dict[str, Any]]:
    """Tier 2 — extract title from the most recent current work-experience entry."""
    current_roles = [e for e in work_experience if _is_current_role(e)]
    if not current_roles:
        current_roles = work_experience[:1]  # best-guess: first entry
    if not current_roles:
        return None

    def _start_year(e: dict) -> int:
        sd = e.get("work_experience_start_date") or {}
        return sd.get("year") or 0

    best  = max(current_roles, key=_start_year)
    title = (best.get("work_experience_job_title") or "").strip()
    if not title:
        return None

    return {
        "current_position":          title,
        "years_of_experience":       None,
        "years_of_experience_value": None,
        "source":                    "work_experience",
    }


def extract_current_position_and_experience(
    sections: dict[str, str],
    work_experience: Optional[list[dict[str, Any]]] = None,
) -> dict[str, Any]:
    """
    Extract the candidate's current / most recent position title and the
    years of experience stated alongside it.

    Priority
    --------
    Tier 1 — summary / about / profile / career_summary section.
              Scans first 15 lines for the pattern:
              "[Adj] [Position] with [X]+ years of experience"
    Tier 2 — most recent work_experience entry with is_current end_date
              or no end_date at all.
    Tier 3 — structured "not identifiable" fallback.

    Returns
    -------
    {
        "current_position":          str,         # e.g. "Senior Software Engineer"
        "years_of_experience":       str | None,  # e.g. "8+ years"
        "years_of_experience_value": float | None,# e.g. 8.0
        "source":                    str          # "summary" | "work_experience" | "none"
    }
    """
    # Tier 1 — try each summary-like section key in priority order
    summary_keys = (
        "about", "summary", "objective", "profile",
        "professional_summary", "career_summary", "others",
    )
    for key in summary_keys:
        summary_text = sections.get(key, "").strip()
        if summary_text:
            result = _scan_summary_for_position(summary_text)
            if result:
                return result

    # Tier 2 — fall back to work experience entries
    if work_experience:
        result = _extract_from_work_experience(work_experience)
        if result:
            return result

    # Tier 3 — not identifiable
    return {
        "current_position":          "Not clearly identifiable",
        "years_of_experience":       None,
        "years_of_experience_value": None,
        "source":                    "none",
    }


# ══════════════════════════════════════════════════════════════════════════════
# ── END PATCH ─────────────────────────────────────────────────────────────────
# ══════════════════════════════════════════════════════════════════════════════


# ── Education parsing ─────────────────────────────────────────────────────────

_INDIAN_STATES = {
    "andhra pradesh", "arunachal pradesh", "assam", "bihar", "chhattisgarh",
    "goa", "gujarat", "haryana", "himachal pradesh", "jharkhand", "karnataka",
    "kerala", "madhya pradesh", "maharashtra", "manipur", "meghalaya",
    "mizoram", "nagaland", "odisha", "punjab", "rajasthan", "sikkim",
    "tamil nadu", "telangana", "tripura", "uttar pradesh", "uttarakhand",
    "west bengal", "delhi", "chandigarh", "pondicherry", "puducherry",
    "jammu", "kashmir", "tamilnadu", "tamilnad",
}
_SCHOOL_LEVELS = {"12th", "10th"}
_ORG_LOCATION_SUFFIX_RE = re.compile(
    r",?\s*(?:" + "|".join(re.escape(s) for s in sorted(_INDIAN_STATES, key=len, reverse=True))
    + r"|india|IN)\s*[,.]?\s*(?:india|IN)?\s*$", re.IGNORECASE,
)
_ORG_TRAILING_YEAR_RE = re.compile(r"\s+(?:19|20)\d{2}(?:\s+[\d.]+(?:/\d+)?)?\s*$")
_ORG_MONTH_YEAR_RE    = re.compile(
    r"\s*,?\s*(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w*\s+(?:19|20)\d{2}.*$",
    re.IGNORECASE,
)
_ORG_GRADE_SUFFIX_RE  = re.compile(
    r"\s*[-–,]?\s*(?:S\.?S\.?C|HSC|SSLC|10th|12th|CGPA|GPA|Percentage|%)\b.*$",
    re.IGNORECASE,
)
_TABLE_HEADER_RE = re.compile(
    r"institution\s*(?:board|/|year)|board\s*/\s*year|year\s*marks|"
    r"degree\s*university\s*year|qualification\s*board|marksin|marks\s*in\s*%",
    re.IGNORECASE,
)
_WORK_DESC_RE = re.compile(
    r"^(?:project\s+description|responsibility|environment\s*:|client\s+name|"
    r"company\s+name|maintain(?:ing)?|support\s+and|rfid|boom\s+barrier|"
    r"dynamic\s+report|daily\s+backup|employee\s+attendance|billing\s+software|"
    r"payroll|student\s+attendance|check\s+daily)",
    re.IGNORECASE,
)
_EDU_NOISE_RE = re.compile(
    r"^(?:education(?:al)?(?:\s+(?:background|qualifications?|details|history))?|"
    r"academic\s+(?:profile|credentials?|qualification|background|details|history)|"
    r"qualifications?|course\s+institution|board\s*/\s*year|"
    r"relevant\s+coursework|level\s*:\s*graduate|marks\s*in|"
    r"working\s+experience|work\s+experience)\s*$",
    re.IGNORECASE,
)
_JUNK_LINE_RE = re.compile(
    r"(?:https?://|www\.|linkedin\.com|github\.com|mailto:|@\w+\.\w+"
    r"|^\s*[*•✓➔→◦·▪▸◆●¢«]\s|\bLeetCode\b|\bHackerRank\b"
    r"|\bAward\b.*\bAwarded\b|\bSolved\s+\d+\+|\bCertif(?:ied|ication)\b"
    r"|^Relevant\s+Coursework|^(?:Coursework|Courses)\s*:)",
    re.IGNORECASE,
)
_DEG_IS_ORG_RE = re.compile(
    r"\b(?:college|university|institute|school|polytechnic|academy|vidyapeeth"
    r"|campus|iit|nit|iiit|iim|ignou|autonomous)\b",
    re.IGNORECASE,
)
_PLACE_NAME_RE = re.compile(
    r"^(?:chennai|mumbai|pune|delhi|bangalore|bengaluru|hyderabad|coimbatore|"
    r"madurai|trichy|tiruchirappalli|salem|erode|tirupur|tirunelveli|vellore|"
    r"kolkata|ahmedabad|surat|jaipur|lucknow|kanpur|nagpur|indore|thane|bhopal|"
    r"visakhapatnam|vadodara|aurangabad|nashik|rajkot|meerut|tamilnadu|tamilnad|"
    r"maharashtra|karnataka|kerala|andhra\s+pradesh|raigad|ahmednagar|shirpur|"
    r"dindigul|(?:uttar|madhya)\s+pradesh)\s*(?:,|$)",
    re.IGNORECASE,
)
_ORG_PREFIX_NOISE_RE  = re.compile(r"^[eo]\s+(?=[A-Z])", re.UNICODE)
_INVALID_FOS_RE = re.compile(
    r"(?:skill|technical|exposure|tool|language|soft|hard|competenc|framework|"
    r"aggregat|section|department|contact|detail|personal|profile|summary|"
    r"experience|working|responsib)",
    re.IGNORECASE,
)
_AFFILIATING_UNIV_RE = re.compile(
    r"\b(?:anna\s+university|thiruvalluvar\s+university|madurai\s+kamaraj\s+university|"
    r"bharathiar\s+university|bharathidasan\s+university|manonmaniam\s+sundaranar|"
    r"periyar\s+university|alagappa\s+university|annamalai\s+university|"
    r"mother\s+teresa\s+university|sathyabama\s+university|vtu\b|"
    r"mumbai\s+university|pune\s+university|osmania\s+university|"
    r"rajasthan\s+university|delhi\s+university|calcutta\s+university|"
    r"ignou\b|autonomous)\b",
    re.IGNORECASE,
)
_WORK_INLINE_GUARD_RE = re.compile(
    r"(?:manager|engineer|developer|analyst|executive|officer|lead|intern|"
    r"architect|consultant|specialist|coordinator|director|head\s+of|sr\.|jr\.)",
    re.IGNORECASE,
)
_KV_YEAR_RE  = re.compile(r"(?:year\s+of\s+pass(?:ing)?|year(?:\s+of)?)\s*[:\-]\s*([\d\s\-–]+)", re.IGNORECASE)
_KV_GRADE_RE = re.compile(r"(?:percentage|cgpa|gpa|grade|marks|score)\s*[:\-]\s*([\d.]+\s*%?)", re.IGNORECASE)
_KV_INST_RE  = re.compile(r"institution\s*[:\-]\s*(.+)", re.IGNORECASE)
_KV_DEG_RE   = re.compile(r"(?:ug\s+degree|pg\s+degree|degree)\s*[:\-]\s*(.+)", re.IGNORECASE)

_DEGREE_LABEL_MAP: list[tuple[str, str, str]] = [
    (r"\bPh\.?D\b|\bDoctorate\b|\bD\.Phil\b|\bpost.?doc\b",    "Ph.D",        "doctoral"),
    (r"\bM\.?Tech\b",                                            "M.Tech",      "master"),
    (r"\bM\.?E\b(?:\s*[-–(]|\s+in\b)",                          "M.E",         "master"),
    (r"\bMBA\b",                                                 "MBA",         "master"),
    (r"\bMCA\b|\bM\.?C\.?A\b",                                  "MCA",         "master"),
    (r"\bM\.?Sc\b",                                              "M.Sc",        "master"),
    (r"\bM\.?Com\b",                                             "M.Com",       "master"),
    (r"\bM\.?Phil\b",                                            "M.Phil",      "master"),
    (r"\bPGDM\b|\bPGD\b",                                       "PGDM",        "master"),
    (r"\bPGDIBO\b",                                              "PGDIBO",      "master"),
    (r"\bPost\s+Graduate\s+Diploma\b",                          "PG Diploma",  "master"),
    (r"\bMaster(?:s|\'s)?\s+(?:of|in)\b",                      "Master",      "master"),
    (r"\bB\.?Tech\b|\bBachelor\s+of\s+Technology\b",           "B.Tech",      "bachelor"),
    (r"\bB\.?E\b(?:\s*[-–(]|\s+in\b|\s+\()|"
     r"\bBachelor\s+of\s+Engineering\b",                        "B.E",         "bachelor"),
    (r"\bBCA\b",                                                "BCA",         "bachelor"),
    (r"\bBCS\b|\bBachelor\s+of\s+Computer\b",                  "BCS",         "bachelor"),
    (r"\bB\.?Sc\b|\bBSC\b",                                    "B.Sc",        "bachelor"),
    (r"\bB\.?Com\b|\bBCOM\b",                                  "B.Com",       "bachelor"),
    (r"\bBBA\b",                                                "BBA",         "bachelor"),
    (r"\bBMS\b",                                                "BMS",         "bachelor"),
    (r"\bMBBS\b",                                               "MBBS",        "bachelor"),
    (r"\bBachelor(?:s|\'s)?\s+(?:of|in)\b",                    "Bachelor",    "bachelor"),
    (r"\bB\s*E\s+Information\b",                                "B.E",         "bachelor"),
    (r"\bAdvanced\s+Diploma\b",                                 "Adv. Diploma","diploma"),
    (r"\bDiploma\b|\bPolytechnic\b|\bITI\b",                   "Diploma",     "diploma"),
    (r"\bHSC\b|\bClass\s*XII\b|\b12th\b|\bHigher\s+Secondary\b"
     r"|\bIntermediate\b|\bPUC\b|\bSenior\s+Secondary\b",      "HSC",         "12th"),
    (r"\bSSLC\b|\bS\.S\.C\b|\bSSC\b|\bClass\s*X\b|\b10th\b"
     r"|\bMatriculation\b|\bHigh\s+School\b|"
     r"\bSecondary\s+School\s+Leaving\b",                       "SSLC",        "10th"),
]

_FIELD_MAP: dict[str, list[str]] = {
    "Computer Science and Engineering":  ["computer science and engineering", r"\bcse\b"],
    "Computer Science":                  ["computer science"],
    "Information Technology":            ["information technology"],
    "Electronics and Communication":     ["electronics and communication", r"\bece\b", r"e\s*&\s*c"],
    "Electrical and Electronics":        [r"electrical.*electronics", r"\beee\b"],
    "Electrical Engineering":            [r"\belectrical engineering\b"],
    "Mechanical Engineering":            [r"\bmechanical engineering\b"],
    "Computer Applications":             ["computer applications"],
    "Mathematics":                       [r"\bmathematics\b"],
    "Physics":                           [r"\bphysics\b"],
    "IT & Management":                   [r"it\s*&\s*m"],
    "Commerce":                          [r"\bcommerce\b"],
    "Business Administration":           ["business administration"],
    "Project Management":                ["project management"],
    "Civil Engineering":                 [r"\bcivil engineering\b"],
    "Biotechnology":                     [r"\bbiotechnology\b"],
    "Artificial Intelligence and ML":    ["artificial intelligence", "machine learning", r"\baiml\b", r"\bai\b", r"\bml\b"],
    "Data Science":                      ["data science", r"\bds\b", "data analytics", "big data"],
    "Chemical Engineering":              ["chemical engineering", r"\bch\b", r"\bchem\b"],
    "Aerospace Engineering":             ["aerospace engineering", r"\baero\b"],
    "Electronics Engineering":           ["electronics engineering", r"\belectronics\b", r"\beie\b"],
    "Software Engineering":              ["software engineering", r"\bse\b"],
    "Economics":                         ["economics"],
    "English Literature":                ["english literature", "english"],
    "Medicine":                          ["mbbs", "medicine", "medical"],
    "Pharmacy":                          ["pharmacy", r"\bbpharm\b", r"\bdpharm\b"],
    "Nursing":                           ["nursing", r"\bbsc nursing\b"],
}


def _detect_education_level(line: str) -> str:
    lower = line.lower()
    for level, keywords in EDUCATION_LEVELS.items():
        if any(kw in lower for kw in keywords):
            return level
    return "other"


def _normalise_degree(text: str) -> tuple[str, str]:
    for pattern, label, level in _DEGREE_LABEL_MAP:
        if re.search(pattern, text, re.IGNORECASE):
            return label, level
    return text.strip(), _detect_education_level(text)


def _extract_field_of_study(text: str) -> Optional[str]:
    for field_name, patterns in _FIELD_MAP.items():
        for p in patterns:
            if re.search(p, text, re.IGNORECASE):
                return field_name
    m = re.search(r"\b(?:in|of)\s+([A-Za-z &/()\-]+)", text, re.IGNORECASE)
    if m:
        candidate = m.group(1).strip().strip("()")
        if 3 < len(candidate) < 60:
            return candidate
    return None


def _parse_edu_grade(line: str) -> Optional[dict[str, Any]]:
    kv = _KV_GRADE_RE.search(line)
    if kv:
        raw = kv.group(1).replace("%", "").strip()
        try:
            score = float(raw)
        except ValueError:
            return None
        unit = "Percentage" if ("%" in kv.group(1) or score > 10) else "CGPA"
        return {"score": score, "unit": unit, "original": line.strip()}
    m = _SCORE_RE.search(line)
    if not m:
        return None
    raw_val = m.group(1) or m.group(2)
    if not raw_val:
        return None
    try:
        score = float(raw_val)
    except ValueError:
        return None
    denom = m.group(3)
    lower = line.lower()
    if denom == "10" or "cgpa" in lower or "gpa" in lower or score <= 10:
        unit = "CGPA"
    elif denom == "100" or "%" in line or "percentage" in lower or score > 10:
        unit = "Percentage"
    else:
        unit = "Unknown"
    return {"score": score, "unit": unit, "original": line.strip()}


def _extract_years_from_line(line: str) -> tuple[Optional[int], Optional[Any]]:
    kv = _KV_YEAR_RE.search(line)
    if kv:
        raw = kv.group(1).strip()
        yrs = re.findall(r"\b((?:19|20)\d{2})\b", raw)
        if len(yrs) >= 2: return int(yrs[0]), int(yrs[-1])
        if len(yrs) == 1: return None, int(yrs[0])
    if re.search(r"\b(?:Pursuing|Distance)\b", line, re.IGNORECASE):
        ym = re.search(r"\b((?:19|20)\d{2})\b", line)
        if ym: return None, "Pursuing"
    pm = _PAREN_YEAR_RE.search(line)
    if pm:
        start   = int(pm.group(1))
        end_raw = pm.group(2)
        end: Any = None
        if end_raw:
            end = "Present" if not end_raw.isdigit() else int(end_raw)
        return start, end
    mym = re.search(
        r"(?:[A-Za-z]{3,9}\.?\s+)?((?:19|20)\d{2})\s*[-–to]+\s*"
        r"(?:[A-Za-z]{3,9}\.?\s+)?((?:19|20)\d{2}|[Pp]ursuing|[Pp]resent|[Cc]urrent)\b",
        line, re.IGNORECASE,
    )
    if mym:
        s = int(mym.group(1)); e_raw = mym.group(2)
        return s, (e_raw if not e_raw.isdigit() else int(e_raw))
    ym = re.search(
        r"\b((?:19|20)\d{2})\s*[-–]\s*((?:19|20)\d{2}|[Pp]resent|[Cc]urrent)\b",
        line, re.IGNORECASE,
    )
    if ym:
        s = int(ym.group(1)); e_raw = ym.group(2)
        return s, (e_raw if not e_raw.isdigit() else int(e_raw))
    ym_s = re.search(r"\b((?:19|20)\d{2})\b", line)
    if ym_s: return None, int(ym_s.group(1))
    return None, None


def _is_degree_line(line: str) -> bool:
    lower = line.lower()
    for _lvl, keywords in EDUCATION_LEVELS.items():
        for kw in keywords:
            if len(kw) <= 4:
                if re.search(r"\b" + re.escape(kw) + r"\b", lower): return True
            elif kw in lower: return True
    return False


def _looks_like_institution(line: str) -> bool:
    return any(kw in line.lower() for kw in [
        "college", "university", "institute", "institution", "school",
        "polytechnic", "vidyapeeth", "technologies", "academy", "ignou",
        "campus", "faculty", "iit", "nit", "iiit", "iim", "vidyalaya",
        "deemed", "autonomous",
    ])


def _clean_org_name(org: str) -> str:
    if not org: return org
    org = _ORG_PREFIX_NOISE_RE.sub("", org).strip()
    org = re.sub(r"\s*\n.*$", "", org, flags=re.DOTALL).strip()
    org = _ORG_GRADE_SUFFIX_RE.sub("", org).strip()
    org = _ORG_TRAILING_YEAR_RE.sub("", org).strip()
    org = _ORG_MONTH_YEAR_RE.sub("", org).strip()
    for _ in range(2):
        stripped = _ORG_LOCATION_SUFFIX_RE.sub("", org).strip().rstrip(",").strip()
        if stripped == org: break
        if len(stripped) < 5: break
        if _PLACE_NAME_RE.match(stripped) and not _looks_like_institution(stripped): break
        org = stripped
    return org.strip().rstrip(",").strip()


def _is_valid_degree(deg: str) -> bool:
    if not deg or len(deg.strip()) < 2: return False
    if _JUNK_LINE_RE.search(deg): return False
    if _TABLE_HEADER_RE.search(deg): return False
    if _WORK_DESC_RE.search(deg): return False
    if _DEG_IS_ORG_RE.search(deg): return False
    if (_PLACE_NAME_RE.match(deg.strip()) or re.match(r"^[A-Z][a-z]+,\s*[A-Z]", deg.strip())) and not _is_degree_line(deg):
        return False
    if len(deg) > 100 and not _is_degree_line(deg): return False
    return True


def _org_similarity(a: str, b: str) -> float:
    a_tok = set(re.findall(r"\b\w{3,}\b", a.lower()))
    b_tok = set(re.findall(r"\b\w{3,}\b", b.lower()))
    if not a_tok or not b_tok: return 0.0
    return len(a_tok & b_tok) / max(len(a_tok), len(b_tok))


def _extract_institutions_nlp(text: str) -> list[str]:
    try:
        nlp = get_nlp()
        doc = nlp(text[:5000])
        return [ent.text.strip() for ent in doc.ents if ent.label_ == "ORG" and _looks_like_institution(ent.text)]
    except Exception as exc:
        log.debug("NLP institution extraction failed: %s", exc)
        return []


def _regex_parse_education(text: str) -> list[Education]:
    raw_lines = _lines(text)
    expanded: list[str] = []
    for line in raw_lines:
        if line.count("|") >= 2 or line.count("\t") >= 2:
            sep = "|" if line.count("|") >= line.count("\t") else "\t"
            expanded.extend(p.strip() for p in line.split(sep) if p.strip())
        else:
            expanded.append(line)
    raw_lines = expanded
    raw_lines = [
        ln for ln in raw_lines
        if not _EDU_NOISE_RE.match(ln.strip())
        and not _JUNK_LINE_RE.search(ln)
        and not _TABLE_HEADER_RE.search(ln)
        and not _WORK_DESC_RE.search(ln)
    ]
    merged: list[str] = []
    skip = False
    for idx, line in enumerate(raw_lines):
        if skip:
            skip = False
            continue
        if (_looks_like_institution(line) and not _is_degree_line(line) and idx + 1 < len(raw_lines)):
            nxt = raw_lines[idx + 1]
            if (not _is_degree_line(nxt) and not _looks_like_institution(nxt)
                    and not re.search(r"\b\d{4}\b", nxt) and len(nxt) < 70):
                merged.append(line.rstrip() + " " + nxt.strip())
                skip = True
                continue
        merged.append(line)

    inline_results: list[Education] = []
    inline_consumed: set[int] = set()
    _COMPLETED_FROM_RE = re.compile(r"Completed\s+(.+?)\s+from\s+(.+?)\s*\((\d{4})\)", re.IGNORECASE)
    _INLINE_DOT_RE = re.compile(r"^(.+?)\s*\((\d{4})\)\s*[.]\s+(.+?)(?:,\s*.+)?$", re.IGNORECASE)
    _DASH_COMMA_INLINE_RE = re.compile(r"^(.+?)\s*[-–]\s*(\d{4})\s*,\s*(.+)", re.IGNORECASE)

    for idx, line in enumerate(merged):
        clean_line = _BULLET_RE.sub("", line).strip()
        m = _COMPLETED_FROM_RE.search(clean_line)
        if m:
            deg_raw, org_raw, yr = m.group(1).strip(), m.group(2).strip(), int(m.group(3))
            label, level = _normalise_degree(deg_raw)
            inline_results.append(Education(organization=org_raw, degree=label, level=level, end_year=yr))
            inline_consumed.add(idx); continue
        m2 = _INLINE_DOT_RE.match(clean_line)
        if m2:
            deg_part, yr, org_part = m2.group(1).strip(), int(m2.group(2)), m2.group(3).strip()
            if _is_degree_line(deg_part) and _looks_like_institution(org_part):
                label, level = _normalise_degree(deg_part)
                field = _extract_field_of_study(deg_part)
                inline_results.append(Education(organization=org_part.split(",")[0].strip(), degree=label, level=level, end_year=yr, field_of_study=field or ""))
                inline_consumed.add(idx); continue
        m3 = _DASH_COMMA_INLINE_RE.match(clean_line)
        if m3:
            deg_part, yr, rest = m3.group(1).strip(), int(m3.group(2)), m3.group(3).strip()
            if _is_degree_line(deg_part) and not _WORK_INLINE_GUARD_RE.search(deg_part):
                deg_clean = re.sub(r"\s*\[.+?\]", "", deg_part).strip()
                label, level = _normalise_degree(deg_clean)
                field = _extract_field_of_study(deg_part)
                parts = [p.strip() for p in rest.split(",")]
                org_primary = parts[0] if parts else ""
                if not _looks_like_institution(org_primary) and level not in _SCHOOL_LEVELS:
                    org_primary = ""
                inline_results.append(Education(organization=org_primary, degree=label, level=level, end_year=yr, field_of_study=field or ""))
                inline_consumed.add(idx); continue

    merged = [ln for i, ln in enumerate(merged) if i not in inline_consumed]
    items: list[Education] = []
    current: Optional[Education] = None

    def _commit() -> None:
        nonlocal current
        if current and (current.organization or current.degree):
            items.append(current)
        current = None

    def _apply_years(edu: Education, line: str) -> None:
        sy, ey = _extract_years_from_line(line)
        if sy and not edu.start_year: edu.start_year = sy
        if ey and not edu.end_year:   edu.end_year   = ey

    def _apply_degree(edu: Education, line: str) -> None:
        clean = line.strip()
        clean = re.sub(
            r"(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+)?"
            r"\b((?:19|20)\d{2})\s*[-–]?\s*"
            r"(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+)?"
            r"((?:19|20)\d{2}|[Pp]ursuing|[Pp]resent|[Cc]urrent)?\b", "", clean,
        )
        clean = re.sub(r"[-–]\s*[\d.]+\s*%", "", clean)
        clean = re.sub(r"[\d.]+\s*/\s*(?:10|100)", "", clean)
        clean = re.sub(r",?\s*(?:CGPA|GPA|Percentage|Marks)\s*:?\s*[\d.]+.*$", "", clean, flags=re.IGNORECASE)
        clean = clean.strip("().,- ")
        if not clean: return
        if _PLACE_NAME_RE.match(clean) and not _is_degree_line(clean): return
        label, level = _normalise_degree(clean)
        edu.degree = label; edu.level = level
        field = _extract_field_of_study(clean)
        if field and not _INVALID_FOS_RE.search(field) and len(field) < 60 and "(" not in field:
            edu.field_of_study = field

    for line in merged:
        if _is_separator(line) or _EDU_NOISE_RE.match(line.strip()): continue
        if _TABLE_HEADER_RE.search(line) or _WORK_DESC_RE.search(line): continue
        kv_inst = _KV_INST_RE.match(line)
        if kv_inst:
            if current: current.organization = kv_inst.group(1).strip(); _apply_years(current, line)
            continue
        kv_deg = _KV_DEG_RE.match(line)
        if kv_deg:
            if current and current.degree: _commit()
            if current is None: current = Education()
            _apply_degree(current, kv_deg.group(1)); _apply_years(current, line); continue
        if _KV_GRADE_RE.match(line) and current:
            grade = _parse_edu_grade(line)
            if grade and not current.grade: current.grade = grade
            continue
        if _KV_YEAR_RE.match(line) and current: _apply_years(current, line); continue
        is_degree   = _is_degree_line(line)
        is_pure_org = _looks_like_institution(line) and not is_degree
        if is_pure_org:
            if current is not None:
                if current.degree and not current.organization:
                    if current.level not in _SCHOOL_LEVELS:
                        current.organization = line.strip(); _apply_years(current, line)
                        grade = _parse_edu_grade(line)
                        if grade and not current.grade: current.grade = grade
                elif current.organization and not current.degree:
                    _commit(); current = Education(organization=line.strip()); _apply_years(current, line)
                elif current.organization and current.degree:
                    _commit(); current = Education(organization=line.strip()); _apply_years(current, line)
                else:
                    current.organization = line.strip(); _apply_years(current, line)
            else:
                current = Education(organization=line.strip()); _apply_years(current, line)
            continue
        if is_degree:
            if line.strip().startswith("(") and current and current.degree:
                frag = line.strip().strip("()")
                if frag and not current.field_of_study: current.field_of_study = frag
                continue
            if current and current.organization and not current.degree:
                _apply_degree(current, line); _apply_years(current, line)
                grade = _parse_edu_grade(line)
                if grade and not current.grade: current.grade = grade
                continue
            if current and current.organization: _commit()
            if current is None: current = Education()
            if current and current.degree and not current.organization: _commit(); current = Education()
            _apply_degree(current, line); _apply_years(current, line)
            grade = _parse_edu_grade(line)
            if grade and current and not current.grade: current.grade = grade
            continue
        if current is None: continue
        pm_m = _PAREN_YEAR_RE.match(line.strip())
        if pm_m and not current.start_year:
            current.start_year = int(pm_m.group(1))
            end_raw = pm_m.group(2)
            if end_raw: current.end_year = "Present" if not end_raw.isdigit() else int(end_raw)
            continue
        grade = _parse_edu_grade(line)
        if grade and not current.grade: current.grade = grade; continue
        sy, ey = _extract_years_from_line(line)
        if ey and not current.end_year:
            if sy and not current.start_year: current.start_year = sy
            current.end_year = ey; continue
        if (not current.organization and _looks_like_institution(line) and current.level not in _SCHOOL_LEVELS
                and len(line) < 120 and not (_PLACE_NAME_RE.match(line.strip()) and not _looks_like_institution(line))):
            current.organization = line.strip(); continue
        if (not current.field_of_study and current.degree and len(line) < 60 and not re.search(r"\d", line)
                and not _is_bullet_line(line) and not _PLACE_NAME_RE.match(line.strip())
                and not _INVALID_FOS_RE.search(line) and "(" not in line):
            current.field_of_study = line.strip()

    _commit()
    return inline_results + [e for e in items if e.organization or e.degree]


def _merge_edu_results(nlp_orgs: list[str], regex_items: list[Education]) -> list[Education]:
    used_nlp: set[int] = set()
    for edu in regex_items:
        if edu.organization or edu.level in _SCHOOL_LEVELS: continue
        for i, org in enumerate(nlp_orgs):
            if i in used_nlp: continue
            if _org_similarity(org, edu.degree or "") > 0.5: continue
            edu.organization = org; used_nlp.add(i); break
    for i, org in enumerate(nlp_orgs):
        if i in used_nlp: continue
        already_covered = any(_org_similarity(org, e.organization or "") > 0.6 for e in regex_items)
        if not already_covered: regex_items.append(Education(organization=org))
    return regex_items


def _post_filter(items: list[Education]) -> list[Education]:
    cleaned: list[Education] = []
    for edu in items:
        if edu.organization: edu.organization = _clean_org_name(edu.organization)
        if edu.degree and _DEG_IS_ORG_RE.search(edu.degree):
            if not edu.organization: edu.organization = _clean_org_name(edu.degree)
            edu.degree = ""; edu.level = ""
        if edu.degree:
            label, level = _normalise_degree(edu.degree)
            edu.degree = label
            if level != "other": edu.level = level
        if edu.degree and not _is_valid_degree(edu.degree):
            edu.degree = ""; edu.level = ""; edu.field_of_study = ""
        if edu.level in _SCHOOL_LEVELS and edu.organization:
            is_pure_place = bool(_PLACE_NAME_RE.match(edu.organization.strip())) and len(edu.organization.strip().split()) <= 3
            if is_pure_place: edu.organization = ""
        if not edu.organization and not edu.degree: continue
        cleaned.append(edu)
    deduped: list[Education] = []
    for edu in cleaned:
        org_a, deg_a, lvl_a = edu.organization or "", edu.degree or "", edu.level or ""
        duplicate = False
        for existing in deduped:
            org_b, deg_b, lvl_b = existing.organization or "", existing.degree or "", existing.level or ""
            org_sim  = _org_similarity(org_a, org_b) if org_a and org_b else 0.0
            deg_match = (deg_a == deg_b) and bool(deg_a)
            lvl_match = (lvl_a == lvl_b) and bool(lvl_a)
            if deg_match and lvl_match:
                both_have_org = bool(org_a) and bool(org_b)
                orgs_differ   = both_have_org and _org_similarity(org_a, org_b) < 0.4
                if orgs_differ: continue
                if edu.start_year   and not existing.start_year:   existing.start_year   = edu.start_year
                if edu.end_year     and not existing.end_year:      existing.end_year     = edu.end_year
                if edu.grade        and not existing.grade:         existing.grade        = edu.grade
                if edu.organization and not existing.organization:  existing.organization = edu.organization
                if edu.field_of_study and not existing.field_of_study: existing.field_of_study = edu.field_of_study
                duplicate = True; break
            if org_sim > 0.7 and (deg_match or not deg_a or not deg_b):
                if len(org_a) > len(org_b): existing.organization = org_a
                if deg_a and not deg_b:
                    existing.degree = deg_a; existing.level = lvl_a or existing.level
                    existing.field_of_study = edu.field_of_study or existing.field_of_study
                if edu.grade      and not existing.grade:      existing.grade      = edu.grade
                if edu.start_year and not existing.start_year: existing.start_year = edu.start_year
                if edu.end_year   and not existing.end_year:   existing.end_year   = edu.end_year
                duplicate = True; break
            if deg_match and not org_a and org_b: duplicate = True; break
            if deg_match and org_a and not org_b: existing.organization = org_a; duplicate = True; break
        if not duplicate: deduped.append(edu)
    result: list[Education] = []
    for edu in deduped:
        is_org_only = not edu.degree and edu.organization and not edu.start_year and not edu.end_year and not edu.grade
        if is_org_only:
            org_lower = (edu.organization or "").lower()
            covered_by_overlap = any(
                (_org_similarity(edu.organization, e.organization or "") > 0.4
                 or org_lower in (e.organization or "").lower()
                 or any(tok in (e.organization or "").lower() for tok in org_lower.split() if len(tok) > 4))
                for e in deduped if e is not edu and (e.degree or e.start_year or e.end_year)
            )
            is_affiliating   = bool(_AFFILIATING_UNIV_RE.search(edu.organization or ""))
            has_real_records = any(e.degree for e in deduped if e is not edu)
            if covered_by_overlap or (is_affiliating and has_real_records): continue
        result.append(edu)
    return result


def parse_education(text: str) -> list[dict[str, Any]]:
    log.debug("parse_education input:\n%s", text)
    if not text.strip():
        return []
    nlp_orgs    = _extract_institutions_nlp(text)
    regex_items = _regex_parse_education(text)
    items       = _merge_edu_results(nlp_orgs, regex_items)
    items       = _post_filter(items)
    items.sort(key=lambda x: x.end_year if isinstance(x.end_year, int) else 9999, reverse=True)
    return [
        {
            "education_organization":   e.organization,
            "education_degree":         e.degree,
            "education_field_of_study": e.field_of_study,
            "education_level":          e.level,
            "education_start_year":     e.start_year,
            "education_end_year":       e.end_year,
            "education_grade":          e.grade,
        }
        for e in items
    ]


# ── Skills parsing ────────────────────────────────────────────────────────────

def _get_skill_category(skill: str) -> str:
    lower = skill.lower()
    for category, keywords in SKILL_CATEGORIES.items():
        if any(kw in lower for kw in keywords):
            return category
    return "Miscellaneous"


def parse_skills(text: str) -> list[str]:
    if not text.strip():
        return []
    try:
        annotations = get_skill_extractor().annotate(text.lower())
        extracted: set[str] = set()
        for match in (
            annotations["results"]["full_matches"]
            + annotations["results"]["ngram_scored"]
        ):
            skill = match["doc_node_value"].strip().title()
            if len(skill) > 2:
                extracted.add(skill)
        return sorted(extracted)
    except Exception as exc:
        log.warning("SkillNer failed (%s) — fallback to regex", exc)
        return _parse_skills_regex(text)


_SKILL_DATE_RE = re.compile(
    r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec).*\d{4}"
    r"|\d{4}\s*[-–]\s*(present|\d{4})",
    re.IGNORECASE,
)
_SKILL_COMPANY_RE = re.compile(
    r"(pvt\.?|ltd\.?|limited|llc|inc\.?|corp\.?)", re.IGNORECASE
)
_SKILL_VERB_RE = re.compile(
    r"^(built|managed|architected|gained|achieved|improved|performed|"
    r"provided|administered|automated|supported|collaborated|actively|"
    r"configured|reduced|reducing|resulting|enabling|leveraging|"
    r"contributing|streamlining|achieving|monitored|handled|implemented|"
    r"developed|created|led|worked|responsible|designed|coordinated|"
    r"participated|ensured|resolved|maintained|deployed|migrated|"
    r"and\s|or\s|the\s|a\s|an\s|in\s|with\s|using\s|via\s|for\s|of\s)",
    re.IGNORECASE,
)
_SKILL_METRIC_RE = re.compile(r"\d+\s*[+%~]|\d{4}")
_SKILL_TITLE_RE = re.compile(
    r"(engineer|developer|analyst|manager|director|officer|lead|intern|"
    r"architect|consultant|specialist|coordinator|trainee|executive|"
    r"researcher|designer|scientist|advisor|head)",
    re.IGNORECASE,
)
_SKILL_JUNK = frozenset({
    "com", "www", "problem", "test", "services", "patching", "restores",
    "backups", "permissions", "prod", "manager", "director", "officer",
    "bengaluru", "mumbai", "delhi", "chennai", "pune", "hyderabad",
    "kolkata", "ahmedabad", "bangalore", "india",
})


def _clean_skill_token(raw: str) -> str:
    token = re.sub(
        r"^(and|or|the|a|an|in|with|using|via|for|of)\s+",
        "", raw, flags=re.IGNORECASE,
    ).strip().strip(".,;:()-")
    if not token or len(token) < 2:
        return ""
    if len(token.split()) > 4:
        return ""
    if _SKILL_DATE_RE.search(token):
        return ""
    if _SKILL_COMPANY_RE.search(token):
        return ""
    if _SKILL_VERB_RE.match(token):
        return ""
    if token.endswith((".", "%", "~")):
        return ""
    if _SKILL_METRIC_RE.search(token):
        return ""
    if _SKILL_TITLE_RE.search(token):
        return ""
    if token.lower() in _SKILL_JUNK:
        return ""
    return token


def _parse_skills_regex(text: str) -> list[str]:
    skills: set[str] = set()
    for line in _lines(text):
        for raw in re.split(r"[,;|•\n]", _clean_bullet(line)):
            token = _clean_skill_token(raw.strip())
            if token:
                skills.add(token)
    return sorted(skills)


def build_skill_objects(names: list[str]) -> list[dict[str, Any]]:
    result = []
    for name in names:
        cat     = _get_skill_category(name)
        is_lang = cat == "Languages (Spoken)"
        result.append({
            "name":        name,
            "category":    cat,
            "type":        "language" if is_lang else "hard",
            "is_language": is_lang,
        })
    return result


def _tokenise(skill: str) -> frozenset[str]:
    tokens = re.findall(r"[a-z0-9]+", skill.lower())
    return frozenset(t for t in tokens if t not in _SKILL_STOPWORDS and len(t) > 1)


def _is_noise(skill: str) -> bool:
    return (
        any(p.search(skill.strip()) for p in _NOISE_PATTERNS)
        or _is_skill_noise_extended(skill)
    )


def _normalise_label(skill: str) -> str:
    key = skill.strip().lower()
    return _KNOWN_FIXES.get(key, skill.strip())


def _collapse_repeated_words(phrase: str) -> str:
    words = phrase.split()
    return " ".join(w for i, w in enumerate(words) if i == 0 or w.lower() != words[i - 1].lower())


def deduplicate_skills(skills: list[str]) -> list[str]:
    cleaned = [s for s in skills if not _is_noise(s)]
    cleaned = [_collapse_repeated_words(s) for s in cleaned]
    cleaned = [_normalise_label(s) for s in cleaned]
    seen_lower: dict[str, str] = {}
    for s in cleaned:
        k = s.lower()
        if k not in seen_lower:
            seen_lower[k] = s
    cleaned = list(seen_lower.values())
    token_groups: dict[frozenset, list[str]] = defaultdict(list)
    for s in cleaned:
        ts = _tokenise(s)
        if ts:
            token_groups[ts].append(s)
    survivors_from_groups: set[str] = set()
    for ts, group in token_groups.items():
        best = max(group, key=len)
        survivors_from_groups.add(best.lower())
    cleaned = [s for s in cleaned if s.lower() in survivors_from_groups]
    final: list[str] = []
    all_token_sets = [(_tokenise(s), s) for s in cleaned]
    for ts_a, s_a in all_token_sets:
        single_token = len(ts_a) == 1
        absorbed = False
        if single_token:
            for ts_b, s_b in all_token_sets:
                if s_b == s_a: continue
                if ts_a < ts_b: absorbed = True; break
        if not absorbed:
            final.append(s_a)
    return final


# ── Projects parsing ──────────────────────────────────────────────────────────

_PROJECT_FIELD_MAP = {
    "client name":  "client",
    "client":       "client",
    "title":        "title",
    "role":         "role",
    "duration":     "duration",
    "technologies": "technologies",
    "tech stack":   "technologies",
    "tools":        "technologies",
    "tech":         "technologies",
    "description":  "description",
}


def parse_projects(text: str) -> list[dict[str, Any]]:
    if not text.strip():
        return []
    lines: list[str] = [_strip_emoji(ln).strip() for ln in _lines(text)]
    items: list[Project] = []
    current: Optional[Project] = None
    for line in lines:
        if _is_separator(line): continue
        if re.match(r"^project\s*\d*\s*[:\-]?", line, re.IGNORECASE):
            if current: items.append(current)
            title   = re.sub(r"^project\s*\d*\s*[:\-]?\s*", "", line, flags=re.IGNORECASE).strip()
            current = Project(title=title)
            continue
        if current is None: current = Project()
        if ":" in line:
            key_raw, _, value = line.partition(":")
            key = key_raw.strip().lower(); value = value.strip()
            mapped = _PROJECT_FIELD_MAP.get(key)
            if mapped == "technologies":
                current.technologies = [t.strip() for t in re.split(r"[,|;]", value) if t.strip()]
                continue
            elif mapped and hasattr(current, mapped):
                setattr(current, mapped, value); continue
        desc_line = _clean_bullet(line)
        if desc_line:
            current.description = (
                (current.description + "\n" + desc_line).strip()
                if current.description else desc_line
            )
    if current: items.append(current)
    return [
        {
            "project_title":        p.title,
            "project_client":       p.client,
            "project_role":         p.role,
            "project_description":  p.description,
            "project_technologies": p.technologies,
            "project_type":         p.project_type,
            **( {"project_dates": p.dates} if p.dates else {} ),
        }
        for p in items
    ]


# ── Certifications parsing ────────────────────────────────────────────────────

def parse_certifications(text: str) -> list[dict[str, str]]:
    items = []
    for line in _lines(text):
        line = _clean_bullet(_strip_emoji(line))
        if not line or _is_separator(line): continue
        entry: dict[str, str] = {"name": line}
        dm = re.search(
            r"\b(20\d{2}|Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s*\d{0,4}",
            line, re.IGNORECASE,
        )
        if dm: entry["date"] = dm.group(0).strip()
        im = re.search(r"(?:by|from|–|-)\s+(.+)$", line, re.IGNORECASE)
        if im: entry["issuer"] = im.group(1).strip()
        items.append(entry)
    return items


# ── Languages parsing ─────────────────────────────────────────────────────────

_PROFICIENCY_MAP = {
    "native":         "Native or bilingual proficiency",
    "bilingual":      "Native or bilingual proficiency",
    "fluent":         "Full professional proficiency",
    "professional":   "Professional working proficiency",
    "conversational": "Limited working proficiency",
    "basic":          "Elementary proficiency",
    "beginner":       "Elementary proficiency",
}


def parse_languages(text: str) -> list[dict[str, Any]]:
    items = []
    for line in _lines(text):
        line = _clean_bullet(line).strip()
        if not line or _is_separator(line): continue
        proficiency  = "Native or bilingual proficiency"
        matched_lang = line
        for key, val in _PROFICIENCY_MAP.items():
            if key in line.lower():
                proficiency  = val
                matched_lang = re.sub(key, "", line, flags=re.IGNORECASE).strip(" -–:,")
                break
        lang = re.split(r"[-–(|]", matched_lang)[0].strip()
        if lang and len(lang) > 1:
            items.append({
                "language_name":        {"label": lang, "value": lang},
                "language_proficiency": {"value": proficiency},
            })
    return items


# ── Generic list parser ───────────────────────────────────────────────────────

def parse_line_list(text: str) -> list[str]:
    if not text.strip():
        return []
    items: list[str] = []
    current_item: str = ""
    for line in _lines(text):
        clean = _strip_emoji(line).strip()
        if _is_separator(clean): continue
        stripped = _clean_bullet(clean)
        if not stripped: continue
        if _is_bullet_line(clean):
            if current_item: items.append(current_item.strip())
            current_item = stripped
        else:
            current_item = (current_item + " " + stripped) if current_item else stripped
    if current_item: items.append(current_item.strip())
    return [i for i in items if i]


# ── Location extraction ───────────────────────────────────────────────────────

_HTTP_USER_AGENT         = "resume-parser/2.0"
_LOCATION_HTTP_TIMEOUT_SEC = 3.0
_LOCATION_EXECUTOR = ThreadPoolExecutor(
    max_workers=2,
    thread_name_prefix="location_http",
)


def _lookup_pincode(pincode: str) -> dict[str, str]:
    try:
        resp = requests.get(
            f"https://api.postalpincode.in/pincode/{pincode}",
            timeout=8, headers={"User-Agent": _HTTP_USER_AGENT},
        )
        data = resp.json()
        if not data or data[0].get("Status") != "Success":
            return {}
        po = (data[0].get("PostOffice") or [{}])[0]
        return {
            "city":         po.get("Block") or po.get("Name") or po.get("District"),
            "district":     po.get("District"),
            "state":        po.get("State"),
            "country":      "India",
            "country_code": "IN",
        }
    except Exception as exc:
        log.debug("Pincode lookup failed: %s", exc)
        return {}


def _geocode_address(addr: str) -> dict[str, str]:
    try:
        url = (
            "https://nominatim.openstreetmap.org/search"
            f"?q={requests.utils.quote(addr + ', India')}"
            "&format=json&limit=1&addressdetails=1"
        )
        resp = requests.get(url, timeout=10, headers={"User-Agent": _HTTP_USER_AGENT})
        data = resp.json()
        if data:
            a = data[0].get("address", {})
            return {
                "formatted":    data[0].get("display_name"),
                "city":         a.get("city") or a.get("town") or a.get("village"),
                "state":        a.get("state"),
                "district":     a.get("county") or a.get("district"),
                "country":      a.get("country", "India"),
                "country_code": (a.get("country_code") or "IN").upper(),
                "postal_code":  a.get("postcode"),
            }
    except Exception as exc:
        log.debug("Geocode failed: %s", exc)
    return {}


def extract_location(resume_text: str, sections: dict[str, str]) -> dict[str, Any]:
    location: dict[str, Any] = {}
    pm = _PINCODE_RE.search(resume_text)
    if pm:
        pincode = pm.group(0)
        location["postal_code"] = pincode
        try:
            fut      = _LOCATION_EXECUTOR.submit(_lookup_pincode, pincode)
            pin_data = fut.result(timeout=_LOCATION_HTTP_TIMEOUT_SEC)
        except FutureTimeoutError:
            log.warning("Pincode lookup timed out for %s", pincode)
            pin_data = {}
        except Exception as exc:
            log.warning("Pincode lookup error for %s: %s", pincode, exc)
            pin_data = {}
        if pin_data:
            location.update(pin_data)
            return {k: v for k, v in location.items() if v}
    addr_m = re.search(
        r"(?:Address|Addr\.?|Location|Residence|Permanent Address|Current Address)"
        r"\s*[:\-]\s*(.+?)(?=\n[A-Z ]{3,}|\n\n|$)",
        resume_text, re.I | re.DOTALL,
    )
    if addr_m:
        raw = addr_m.group(1).strip().replace("\n", " ")
        location["raw_address"] = raw
        try:
            fut = _LOCATION_EXECUTOR.submit(_geocode_address, raw)
            geo = fut.result(timeout=_LOCATION_HTTP_TIMEOUT_SEC)
        except FutureTimeoutError:
            log.warning("Geocode timed out for address")
            geo = {}
        except Exception as exc:
            log.warning("Geocode error: %s", exc)
            geo = {}
        if geo: location.update(geo)
        return {k: v for k, v in location.items() if v}
    search_text = " ".join([
        sections.get("personal_details", ""),
        sections.get("about", ""),
        resume_text[:1000],
    ])
    doc = get_nlp()(search_text)
    for ent in doc.ents:
        if ent.label_ == "GPE":
            if ent.text in _STATES and "state" not in location:
                location["state"] = ent.text
            elif _gc.get_cities_by_name(ent.text) and "city" not in location:
                location["city"] = ent.text
    if location:
        location.setdefault("country",      "India")
        location.setdefault("country_code", "IN")
    return {k: v for k, v in location.items() if v}


# ── Summary parsing ───────────────────────────────────────────────────────────

def _parse_summary_to_sentences(summary_text: str) -> Optional[list[str]]:
    if not summary_text or not summary_text.strip():
        return None
    raw  = summary_text.strip()
    text = raw.replace("\r", "").replace("\t", " ")
    cleaned_lines: list[str] = []
    for line in text.splitlines():
        line = re.sub(r"^\s*(?:\d+\.\s+|[•\-\*✓➔→◦·▪▸◆●]+\s*)", "", line).strip()
        if line: cleaned_lines.append(line)
    text = " ".join(cleaned_lines)
    _CONN = r"of|in|on|at|to|and|with|the|for|by|is|are|was|has|have"
    def _fix_merge(m: re.Match) -> str:
        full = m.group(0)
        if full.lower() in {w.lower() for w in _REAL_WORDS_ENDING_WITH_CONNECTIVE}:
            return full
        return m.group(1) + " " + m.group(2)
    text = re.sub(rf"([a-z]{{4,}})({_CONN})(?=[^a-z]|$)", _fix_merge, text)
    text = re.sub(r" {2,}", " ", text).strip()
    protected = text
    for placeholder, keyword in _SUMMARY_PROTECT:
        pattern = r"\b" + re.escape(keyword) if keyword.endswith(".") and len(keyword) <= 5 else re.escape(keyword)
        protected = re.sub(pattern, placeholder, protected, flags=re.IGNORECASE)
    _version_map: dict[str, str] = {}
    def _protect_version(m: re.Match) -> str:
        tok = f"__VER{len(_version_map):04d}__"; _version_map[tok] = m.group(0); return tok
    protected = re.sub(r"\b\d+\.\d+(?:\.\d+)*\b", _protect_version, protected)
    _decimal_map: dict[str, str] = {}
    def _protect_decimal(m: re.Match) -> str:
        tok = f"__DEC{len(_decimal_map):04d}__"; _decimal_map[tok] = m.group(0); return tok
    protected = re.sub(r"\b\d+\.\d+\b", _protect_decimal, protected)
    raw_parts = re.split(r"(?<=[.!?])\s+", protected)
    def _restore(s: str) -> str:
        for tok, orig in _version_map.items(): s = s.replace(tok, orig)
        for tok, orig in _decimal_map.items(): s = s.replace(tok, orig)
        for placeholder, keyword in _SUMMARY_PROTECT: s = s.replace(placeholder, keyword)
        return s
    sentences = [_restore(part).strip() for part in raw_parts]
    sentences = [s for s in sentences if len(s) >= 15]
    return sentences if sentences else None


# ── Core builder ──────────────────────────────────────────────────────────────

def build_structured_resume(resume_text: str) -> dict[str, Any]:
    """Orchestrate all field parsers and return the structured resume dict."""
    sections = segment_resume(resume_text)
    log.info("Sections detected: %s", list(sections.keys()))

    resume: dict[str, Any] = {}
    resume["candidate_name"] = parse_candidate_name(resume_text)
    contact = parse_contact(resume_text)

    resume["email"] = [contact["email"]] if contact["email"] else None
    if contact.get("all_emails"):
        resume["all_emails"] = contact["all_emails"]

    resume["phone_number"] = [{"raw": p, "e164": p} for p in contact["phone"]] or None

    websites: list[dict[str, str]] = []
    if contact["linkedin"]:
        websites.append({"url": contact["linkedin"], "domain": "linkedin.com"})
    if contact["github"]:
        websites.append({"url": contact["github"],   "domain": "github.com"})
    for url in contact["other_urls"]:
        domain = re.sub(r"https?://(?:www\.)?", "", url).split("/")[0]
        websites.append({"url": url, "domain": domain})
    resume["websites"] = websites or None
    resume["location"] = extract_location(resume_text, sections) or None

    dob_m = _DOB_RE.search(resume_text)
    if dob_m:
        try:
            raw_dob = dob_m.group(1).replace(".", "/").replace("-", "/")
            d, m, y = raw_dob.split("/")
            resume["date_of_birth"] = date(int(y), int(m), int(d)).isoformat()
        except Exception:
            pass

    gm = re.search(r"(?:gender|sex)\s*[:\-]\s*(male|female|other)", resume_text, re.I)
    if gm: resume["gender"] = gm.group(1).title()

    nm = re.search(r"(?:nationality|citizenship)\s*[:\-]\s*([A-Za-z]+)", resume_text, re.I)
    if nm: resume["nationality"] = nm.group(1).title()

    resume["summary"] = _parse_summary_to_sentences(sections.get("about", "").strip())

    # ── Skills — three-tier extraction ────────────────────────────────────────
    dense_skill_text = " ".join(filter(None, [
        sections.get("skills",        ""),
        sections.get("tools_breadth", ""),
    ]))
    skillner_skills = parse_skills(dense_skill_text)

    narrative_text = " ".join(filter(None, [
        sections.get("experience", ""),
        sections.get("projects",   ""),
    ]))
    # regex_skills = _parse_skills_regex(narrative_text)
    # ONLY TRUST structured sections
    dense_skill_text = " ".join(filter(None, [
        sections.get("skills", ""),
        sections.get("tools_breadth", "")
    ]))

    skillner_skills = parse_skills(dense_skill_text)

    # optional: limited regex ONLY from skills section
    regex_skills = _parse_skills_regex(sections.get("skills", ""))

    raw_skills = skillner_skills + regex_skills

    cert_skills  = _parse_skills_regex(sections.get("certifications", ""))
    raw_skills   = skillner_skills + regex_skills + cert_skills
    clean_skills = deduplicate_skills(raw_skills)
    resume["skills"] = build_skill_objects(clean_skills)

    exp_text    = sections.get("experience",  "")
    intern_text = sections.get("internships", "")
    resume["work_experience"]        = parse_experience(exp_text)    or None
    resume["internships"]            = parse_experience(intern_text) or None

    _summary_text = sections.get("about", "")
    resume["total_years_experience"] = calculate_total_experience(
        "\n".join(filter(None, [exp_text, intern_text])),
        summary_text=_summary_text,
    )

    # ── NEW: current position + years of experience ───────────────────────────
    resume["current_position"] = extract_current_position_and_experience(
        sections=sections,
        work_experience=resume.get("work_experience"),
    )

    resume["education"]      = parse_education(sections.get("education",      "")) or None
    resume["projects"]       = parse_projects (sections.get("projects",       "")) or None
    resume["certifications"] = parse_certifications(sections.get("certifications", "")) or None
    resume["publications"]   = parse_line_list(sections.get("publications",      "")) or None
    resume["languages"]      = parse_languages (sections.get("languages_spoken", "")) or None
    resume["achievements"]   = parse_line_list(sections.get("achievements",      "")) or None
    resume["volunteer"]      = parse_line_list(sections.get("volunteer",         "")) or None
    resume["interests"]      = parse_line_list(sections.get("interests",         "")) or None
    resume["references"]     = parse_line_list(sections.get("references",        "")) or None
    resume["_sections_detected"] = list(sections.keys())

    return {k: v for k, v in resume.items() if v is not None}


# ── Pipeline entry-point ──────────────────────────────────────────────────────

def process_resume(file_path: Path) -> dict[str, Any]:
    from resume_parser import parse_resume
    log.info("Processing: %s", file_path.name)
    raw_text = parse_resume(file_path)
    return build_structured_resume(raw_text)


# ── CLI entry-point ───────────────────────────────────────────────────────────

RESUME_INPUT_DIR  = "pdf_file"
JSON_OUTPUT_DIR   = "json_resume"
SUPPORTED_FORMATS = {".pdf", ".docx", ".doc"}


def main() -> None:
    resume_dir = Path(RESUME_INPUT_DIR)
    json_dir   = Path(JSON_OUTPUT_DIR)
    json_dir.mkdir(exist_ok=True)

    for file in resume_dir.glob("*.*"):
        if file.suffix.lower() not in SUPPORTED_FORMATS:
            continue
        try:
            print(f"Processing: {file.name}")
            data      = process_resume(file)
            json_path = json_dir / f"{file.stem}.json"
            with open(json_path, "w", encoding="utf-8") as fh:
                json.dump(data, fh, indent=2, ensure_ascii=False)
            print(f"Saved:      {json_path.name}\n")
        except Exception as e:
            traceback.print_exc()
            print(f"Failed:     {file.name} → {type(e).__name__}: {e}\n")


if __name__ == "__main__":
    main()